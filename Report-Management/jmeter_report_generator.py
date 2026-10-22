#!/usr/bin/env python3
"""
JMeter Word Report Generator (parameter-driven).

Membaca hasil JMeter (.jtl, format CSV) dan sebuah file parameter Markdown (.md)
untuk menghasilkan laporan Word yang rapi berisi metrik, temuan otomatis, dan
rekomendasi tindakan (Database, Application/Connection Pool, Infrastructure).

Contoh:
    python jmeter_report_generator.py --jtl result.jtl --params report_parameters.md
    python jmeter_report_generator.py --jtl result.jtl --params report_parameters.md --no-llm
"""

import argparse
import json
import os
import re
import shutil
import tempfile
from datetime import datetime

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import requests
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

MONTHS_ID = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
             "Agustus", "September", "Oktober", "November", "Desember"]

NAVY = RGBColor(26, 54, 93)
DARK_GRAY = RGBColor(45, 55, 72)
GREEN = RGBColor(39, 110, 80)
RED = RGBColor(197, 48, 48)
WHITE = RGBColor(255, 255, 255)

SEVERITY_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
SEVERITY_FILL = {"Critical": "FED7D7", "High": "FEEBC8", "Medium": "FEFCBF", "Low": "E2E8F0"}


# --------------------------------------------------------------------------- #
# 1. Parser file parameter Markdown
# --------------------------------------------------------------------------- #
def _coerce(value):
    """Ubah teks nilai menjadi bool/int/float bila memungkinkan."""
    v = value.strip()
    low = v.lower()
    if low in ("true", "yes", "ya"):
        return True
    if low in ("false", "no", "tidak"):
        return False
    if re.fullmatch(r"-?\d+", v):
        return int(v)
    if re.fullmatch(r"-?\d+\.\d+", v):
        return float(v)
    return v


def as_list(value):
    if not value:
        return []
    return [x.strip() for x in str(value).split(",") if x.strip()]


_KV_RE = re.compile(r"^[-*]\s+([A-Za-z_][A-Za-z0-9_]*)\s*:\s*(.*)$")


def _parse_block(lines):
    """Parse sekumpulan baris menjadi {kv, items, text}."""
    kv, items = {}, []
    for line in lines:
        s = line.strip()
        m = _KV_RE.match(s)
        if m:
            kv[m.group(1)] = _coerce(m.group(2))
        elif s.startswith(("- ", "* ")):
            items.append(s[2:].strip())
    return {"kv": kv, "items": items, "text": "\n".join(lines).strip()}


class ReportParameters:
    """Representasi file parameter .md."""

    def __init__(self, path=None):
        self.meta, self.sections, self.playbook = {}, {}, {}
        if path:
            self._load(path)

    def _load(self, path):
        if not os.path.exists(path):
            raise FileNotFoundError(f"File parameter tidak ditemukan: {path}")
        with open(path, encoding="utf-8") as fh:
            text = fh.read()

        fm = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
        if fm:
            for line in fm.group(1).splitlines():
                if ":" in line:
                    k, v = line.split(":", 1)
                    self.meta[k.strip()] = _coerce(v)
            text = text[fm.end():]

        parts = re.split(r"^## +", text, flags=re.M)
        for part in parts[1:]:
            title, _, body = part.partition("\n")
            title = title.strip()
            if title == "Playbook Rekomendasi":
                subs = re.split(r"^### +", body, flags=re.M)
                for sub in subs[1:]:
                    key, _, sbody = sub.partition("\n")
                    self.playbook[key.strip()] = _parse_block(sbody.splitlines())["kv"]
            else:
                self.sections[title] = _parse_block(body.splitlines())

    def kv(self, section, defaults=None):
        merged = dict(defaults or {})
        merged.update(self.sections.get(section, {}).get("kv", {}))
        return merged

    def items(self, section):
        return self.sections.get(section, {}).get("items", [])

    def text(self, section, default=""):
        sec = self.sections.get(section)
        if not sec:
            return default
        # buang paragraf pengantar pertama? tidak: kirim apa adanya.
        return sec["text"] or default


DEFAULT_SLA = {"max_p95_ms": 1000, "max_error_rate_pct": 1.0, "min_tps": 5}
DEFAULT_FILTER = {"target_host": "", "exclude_label_regex": "", "include_label_regex": "",
                  "drop_first_seconds": 0}
DEFAULT_ANALYSIS = {
    "ramp_up_seconds": 30, "baseline_threads": 2, "baseline_min_samples": 3,
    "concurrency_min_samples": 5, "concurrency_slope_warn_ms": 30, "concurrency_corr_warn": 0.7,
    "fixed_overhead_ms": 400, "small_payload_kb": 5, "payload_warn_kb": 50,
    "slow_endpoint_ms": 500, "tail_ratio_warn": 3, "tail_min_p99_ms": 1000,
    "connect_warn_ms": 100, "datatable_label_regex": "Data List", "latency_share_warn_pct": 90,
}
DEFAULT_LLM = {
    "enabled": False, "api_url": "http://localhost:11434/v1/chat/completions", "api_key": "ollama",
    "model": "llama3.2", "temperature": 0.2, "timeout_seconds": 300, "min_points": 4,
    "max_points": 5, "language": "Bahasa Indonesia",
}


class SafeDict(dict):
    def __missing__(self, key):
        return "{" + key + "}"


# --------------------------------------------------------------------------- #
# 2. Analyzer
# --------------------------------------------------------------------------- #
class JMeterReportGenerator:
    def __init__(self, jtl_filepath, params: ReportParameters, llm_override=None):
        self.jtl_filepath = jtl_filepath
        self.params = params
        self.sla = params.kv("SLA", DEFAULT_SLA)
        self.flt = params.kv("Filter Data", DEFAULT_FILTER)
        self.cfg = params.kv("Analisis", DEFAULT_ANALYSIS)
        self.llm = params.kv("LLM", DEFAULT_LLM)
        if llm_override is not None:
            self.llm["enabled"] = llm_override
        self.context = {k: v for k, v in params.kv("Konteks Sistem").items()
                        if str(v).strip() and not str(v).strip().startswith("(isi")}
        self.df = None
        self.dq = []            # catatan kualitas data
        self.excluded_hosts = {}  # host -> {rows, codes}
        self.findings = []

    # ---- loading ---------------------------------------------------------- #
    def load_data(self):
        print(f"[*] Membaca file result: {self.jtl_filepath} ...")
        wanted = {"timeStamp", "elapsed", "label", "responseCode", "responseMessage", "success",
                  "bytes", "allThreads", "Latency", "Connect", "URL", "threadName"}
        df = pd.read_csv(self.jtl_filepath, usecols=lambda c: c in wanted, low_memory=False)
        df = df.rename(columns={"timeStamp": "timestamp", "Latency": "latency", "Connect": "connect"})
        for col in ("elapsed", "timestamp", "bytes", "allThreads", "latency", "connect"):
            if col in df.columns:
                df[col] = pd.to_numeric(df[col], errors="coerce")
        for col in ("bytes", "allThreads", "latency", "connect"):
            if col not in df.columns:
                df[col] = np.nan
        df["success"] = df["success"].astype(str).str.lower().map({"true": True, "false": False}).fillna(False)
        df["responseCode"] = df["responseCode"].astype(str)
        df = df.dropna(subset=["elapsed", "timestamp"]).copy()
        df["host"] = df["URL"].astype(str).str.extract(r"^https?://([^/:]+)", expand=False) if "URL" in df.columns else ""
        raw_total = len(df)
        print(f"[ok] {raw_total:,} baris dimuat.")

        # Filter host target
        target = str(self.flt.get("target_host", "")).strip()
        if target:
            off = df[df["host"] != target]
            if len(off):
                for host, g in off.groupby("host"):
                    self.excluded_hosts[host] = {
                        "rows": len(g),
                        "codes": g["responseCode"].value_counts().to_dict(),
                        "failed": int((~g["success"]).sum()),
                    }
                    self.dq.append(
                        f"{len(g):,} sampel ke host '{host}' (kode respons {', '.join(sorted(g['responseCode'].unique()))}, "
                        f"{int((~g['success']).sum())} gagal) berada di luar target_host '{target}' dan dikeluarkan dari analisis.")
                df = df[df["host"] == target]
        else:
            hosts = df["host"].dropna().unique()
            if len(hosts) > 1:
                self.dq.append(f"Data memuat {len(hosts)} host berbeda ({', '.join(map(str, hosts))}). "
                               f"Isi target_host pada file parameter agar host lain dikeluarkan.")

        inc = str(self.flt.get("include_label_regex", "")).strip()
        exc = str(self.flt.get("exclude_label_regex", "")).strip()
        if inc:
            df = df[df["label"].str.contains(inc, regex=True, na=False)]
        if exc:
            before = len(df)
            df = df[~df["label"].str.contains(exc, regex=True, na=False)]
            if before - len(df):
                self.dq.append(f"{before - len(df):,} sampel dengan label cocok '{exc}' dikeluarkan.")

        df = df.sort_values("timestamp").reset_index(drop=True)
        df["t"] = (df["timestamp"] - df["timestamp"].min()) / 1000.0
        drop = float(self.flt.get("drop_first_seconds", 0) or 0)
        if drop > 0:
            n0 = len(df)
            df = df[df["t"] >= drop]
            self.dq.append(f"{n0 - len(df):,} sampel pada {drop:.0f} detik pertama dikeluarkan (drop_first_seconds).")
        if df.empty:
            raise ValueError("Tidak ada data tersisa setelah filter. Periksa bagian 'Filter Data' pada file parameter.")

        self.df = df
        self.dq.insert(0, f"Total sampel mentah {raw_total:,}; sampel dianalisis {len(df):,}.")
        if df["bytes"].isna().all():
            self.dq.append("Kolom 'bytes' tidak tersedia, analisis payload dilewati.")

    # ---- statistik -------------------------------------------------------- #
    def _stats(self, d):
        n = len(d)
        if n == 0:
            return {}
        failed = int((~d["success"]).sum())
        secs = d["timestamp"].floordiv(1000)
        per_sec = secs.value_counts()
        wall = max((d["timestamp"].max() - d["timestamp"].min()) / 1000.0, 1.0)
        active = max(int(per_sec.size), 1)
        e = d["elapsed"]
        lat = d["latency"].mean()
        return {
            "total_req": n, "failed_req": failed, "error_rate": failed / n * 100,
            "tps": n / active, "peak_tps": float(per_sec.max()), "wall_tps": n / wall,
            "active_sec": active, "wall_sec": wall,
            "avg": e.mean(), "min": e.min(), "max": e.max(),
            "p50": e.quantile(.50), "p90": e.quantile(.90), "p95": e.quantile(.95), "p99": e.quantile(.99),
            "avg_connect": d["connect"].mean(), "avg_latency": lat,
            "latency_share": (lat / e.mean() * 100) if e.mean() and not np.isnan(lat) else np.nan,
            "avg_bytes": d["bytes"].mean(), "total_bytes": d["bytes"].sum(),
            "max_threads": d["allThreads"].max(),
        }

    def compute(self):
        self.g = self._stats(self.df)
        self.eps = {l: self._stats(g) for l, g in self.df.groupby("label", sort=False)}
        # fase
        ru = float(self.cfg["ramp_up_seconds"])
        self.phases = {}
        if ru > 0:
            self.phases["Ramp-up (0-%ds)" % ru] = self._stats(self.df[self.df["t"] < ru])
            steady = self.df[self.df["t"] >= ru]
            if len(steady):
                self.phases["Steady state"] = self._stats(steady)
        # profil konkurensi
        self.conc, self.slope, self.corr = self._concurrency_profile()
        # baseline beban rendah
        self.baseline = self._baseline()

    def _concurrency_profile(self):
        d = self.df.dropna(subset=["allThreads"])
        if d.empty:
            return pd.DataFrame(), 0.0, 0.0
        grp = d.groupby(d["allThreads"].astype(int))["elapsed"]
        t = pd.DataFrame({"samples": grp.size(), "avg": grp.mean(), "p95": grp.quantile(.95)}).reset_index()
        t = t.rename(columns={"allThreads": "vu"})
        t = t[t["samples"] >= int(self.cfg["concurrency_min_samples"])]
        if len(t) >= 3 and t["vu"].nunique() >= 3:
            slope = float(np.polyfit(t["vu"], t["avg"], 1)[0])
            corr = float(np.corrcoef(t["vu"], t["avg"])[0, 1])
        else:
            slope, corr = 0.0, 0.0
        return t, slope, corr

    def _baseline(self):
        th = int(self.cfg["baseline_threads"])
        d = self.df[self.df["allThreads"] <= th]
        res = {}
        for label, g in d.groupby("label"):
            if len(g) >= int(self.cfg["baseline_min_samples"]):
                res[label] = float(g["elapsed"].median())
        return res

    # ---- temuan otomatis ---------------------------------------------------#
    def detect_findings(self):
        c, sla, F = self.cfg, self.sla, []
        short = lambda labels, n=4: ", ".join(labels[:n]) + (f" (+{len(labels) - n} lainnya)" if len(labels) > n else "")

        # a. kualitas data / konfigurasi
        if self.excluded_hosts:
            hosts = ", ".join(self.excluded_hosts)
            rows = sum(v["rows"] for v in self.excluded_hosts.values())
            F.append({"key": "error_konfigurasi", "severity": "Medium",
                      "title": "Request ke host di luar target / error skrip uji",
                      "evidence": f"{rows} sampel ke {hosts} dikeluarkan dari analisis (target_host: {self.flt.get('target_host')}).",
                      "vars": {"hosts": hosts}})

        # b. error aplikasi
        if self.g["error_rate"] > sla["max_error_rate_pct"]:
            F.append({"key": "error_aplikasi", "severity": "Critical",
                      "title": "Error rate melebihi SLA",
                      "evidence": f"Error rate {self.g['error_rate']:.2f}% ({self.g['failed_req']:,} gagal) vs SLA {sla['max_error_rate_pct']}%.",
                      "vars": {"error_rate": f"{self.g['error_rate']:.2f}"}})

        # c. antrean konkurensi
        if (not self.conc.empty and self.conc["vu"].nunique() >= 5
                and self.slope >= c["concurrency_slope_warn_ms"] and self.corr >= c["concurrency_corr_warn"]):
            lo, hi = self.conc.iloc[0], self.conc.iloc[-1]
            ph = self.phases
            extra = ""
            if len(ph) == 2:
                ru, st = list(ph.values())
                extra = f" Rata-rata fase ramp-up {ru['avg']:.0f} ms vs steady state {st['avg']:.0f} ms."
            F.append({"key": "antrean_konkurensi", "severity": "High",
                      "title": "Latensi naik linear terhadap konkurensi",
                      "evidence": (f"Avg {lo['avg']:.0f} ms pada {int(lo['vu'])} VU menjadi {hi['avg']:.0f} ms pada {int(hi['vu'])} VU "
                                   f"(slope {self.slope:.0f} ms/VU, korelasi {self.corr:.2f}). "
                                   f"Latency {self.g['latency_share']:.0f}% dari Elapsed, Connect {self.g['avg_connect']:.1f} ms.{extra}"),
                      "vars": {"slope": f"{self.slope:.0f}", "corr": f"{self.corr:.2f}"}})

        # d. overhead tetap pada endpoint kecil
        small = {l for l, s in self.eps.items() if s["avg_bytes"] == s["avg_bytes"] and s["avg_bytes"] < c["small_payload_kb"] * 1024}
        flagged = sorted([l for l in small if self.baseline.get(l, 0) >= c["fixed_overhead_ms"]],
                         key=lambda l: -self.baseline[l])
        if len(flagged) >= 3:
            med = float(np.median([self.baseline[l] for l in flagged]))
            F.append({"key": "overhead_tetap", "severity": "High",
                      "title": "Overhead tetap tinggi pada endpoint berpayload kecil",
                      "evidence": (f"{len(flagged)} dari {len(small)} endpoint kecil (<{c['small_payload_kb']} KB) punya median "
                                   f"≥{c['fixed_overhead_ms']} ms saat ≤{c['baseline_threads']} VU (median gabungan {med:.0f} ms): {short(flagged)}."),
                      "vars": {"labels": short(flagged), "baseline_ms": f"{med:.0f}"}})

        # e. payload besar
        total_b = self.g["total_bytes"] or 0
        big = sorted([(l, s) for l, s in self.eps.items()
                      if s["avg_bytes"] == s["avg_bytes"] and s["avg_bytes"] >= c["payload_warn_kb"] * 1024],
                     key=lambda x: -x[1]["avg_bytes"])
        if big and total_b:
            share = sum(s["total_bytes"] for _, s in big) / total_b * 100
            top = ", ".join(f"{l} ({s['avg_bytes'] / 1024:.0f} KB, avg {s['avg']:.0f} ms)" for l, s in big[:3])
            F.append({"key": "payload_besar", "severity": "High" if share >= 50 else "Medium",
                      "title": "Payload respons berlebihan",
                      "evidence": f"{len(big)} endpoint ≥{c['payload_warn_kb']} KB menyumbang {share:.0f}% dari total transfer ({total_b / 1e6:.0f} MB): {top}.",
                      "vars": {"labels": short([l for l, _ in big]), "payload_kb": f"{np.mean([s['avg_bytes'] for _, s in big]) / 1024:.0f}",
                               "share": f"{share:.0f}"}})

        # f. datatable lambat
        rx = str(c["datatable_label_regex"])
        dt = sorted([l for l, s in self.eps.items() if rx and re.search(rx, l) and s["avg"] > c["slow_endpoint_ms"]],
                    key=lambda l: -self.eps[l]["avg"])
        if dt:
            F.append({"key": "datatable_query", "severity": "Medium",
                      "title": "Endpoint datatable di atas ambang waktu",
                      "evidence": f"{len(dt)} endpoint datatable avg >{c['slow_endpoint_ms']} ms: " +
                                  ", ".join(f"{l} ({self.eps[l]['avg']:.0f} ms)" for l in dt[:4]) + ".",
                      "vars": {"labels": short(dt)}})

        # g. tail latency
        tails = sorted([l for l, s in self.eps.items()
                        if s["p99"] > s["p50"] * c["tail_ratio_warn"] and s["p99"] > c["tail_min_p99_ms"]],
                       key=lambda l: -self.eps[l]["p99"])
        if tails:
            t0 = self.eps[tails[0]]
            F.append({"key": "tail_latency", "severity": "Medium",
                      "title": "Lonjakan latensi ekor (P99 >> P50)",
                      "evidence": f"{len(tails)} endpoint dengan P99 >{c['tail_ratio_warn']}x P50 dan >{c['tail_min_p99_ms']} ms; "
                                  f"contoh {tails[0]}: P50 {t0['p50']:.0f} ms, P99 {t0['p99']:.0f} ms.",
                      "vars": {"labels": short(tails)}})

        # h. connect
        if self.g["avg_connect"] > c["connect_warn_ms"]:
            F.append({"key": "network_connect", "severity": "Medium", "title": "Connect Time tinggi",
                      "evidence": f"Connect rata-rata {self.g['avg_connect']:.1f} ms (> {c['connect_warn_ms']} ms).",
                      "vars": {"connect_ms": f"{self.g['avg_connect']:.1f}"}})

        # i. throughput
        if self.g["tps"] < sla["min_tps"]:
            F.append({"key": "throughput_rendah", "severity": "Low", "title": "Throughput di bawah target",
                      "evidence": f"Active TPS {self.g['tps']:.2f} req/s (puncak {self.g['peak_tps']:.0f}) vs target {sla['min_tps']} req/s; "
                                  f"VU maksimum {int(self.g['max_threads']) if self.g['max_threads'] == self.g['max_threads'] else '-'}.",
                      "vars": {"tps": f"{self.g['tps']:.2f}", "min_tps": str(sla["min_tps"])}})

        F.sort(key=lambda f: SEVERITY_ORDER[f["severity"]])
        self.findings = F

    def endpoint_insights(self, s):
        out, c = [], self.cfg
        if s["p95"] > self.sla["max_p95_ms"]:
            out.append(f"P95 {s['p95']:.0f} ms melebihi SLA {self.sla['max_p95_ms']} ms.")
        if s["p99"] > s["p50"] * c["tail_ratio_warn"] and s["p99"] > c["tail_min_p99_ms"]:
            out.append(f"Latensi ekor: P99 {s['p99']:.0f} ms vs P50 {s['p50']:.0f} ms (indikasi cold start, antrean pool, atau lock).")
        if s["error_rate"] > self.sla["max_error_rate_pct"]:
            out.append(f"Error rate {s['error_rate']:.2f}% melebihi batas {self.sla['max_error_rate_pct']}%.")
        if s["avg_connect"] > c["connect_warn_ms"]:
            out.append(f"Connect rata-rata {s['avg_connect']:.1f} ms, indikasi antrean handshake TCP/TLS.")
        if s["avg_bytes"] == s["avg_bytes"] and s["avg_bytes"] >= c["payload_warn_kb"] * 1024:
            out.append(f"Respons besar ({s['avg_bytes'] / 1024:.0f} KB per request).")
        if not out:
            out.append(f"Dalam batas SLA (P95 {s['p95']:.0f} ms, error {s['error_rate']:.2f}%).")
        return out

    # ---- rekomendasi ------------------------------------------------------ #
    def playbook_blocks(self):
        blocks = []
        for f in self.findings:
            pb = self.params.playbook.get(f["key"])
            if not pb:
                continue
            fmt = lambda k: str(pb.get(k, "")).format_map(SafeDict(f["vars"]))
            blocks.append({"severity": f["severity"], "judul": fmt("judul") or f["title"], "evidence": f["evidence"],
                           "rca": fmt("rca"), "database": fmt("database"), "application": fmt("application"),
                           "infrastructure": fmt("infrastructure"), "verifikasi": fmt("verifikasi")})
        return blocks

    def _llm_prompt(self):
        ep = {l: {"avg": round(s["avg"]), "p95": round(s["p95"]), "p99": round(s["p99"]),
                  "error_pct": round(s["error_rate"], 2), "avg_kb": round(s["avg_bytes"] / 1024, 1) if s["avg_bytes"] == s["avg_bytes"] else None,
                  "baseline_ms_low_load": round(self.baseline[l]) if l in self.baseline else None}
              for l, s in self.eps.items()}
        payload = {
            "global": {k: (round(float(v), 2) if isinstance(v, (int, float, np.floating, np.integer)) and v == v else None)
                       for k, v in self.g.items()},
            "sla": self.sla,
            "fase": {k: {"avg": round(v["avg"]), "p95": round(v["p95"]), "tps": round(v["tps"], 2)} for k, v in self.phases.items()},
            "profil_konkurensi": [{"vu": int(r.vu), "samples": int(r.samples), "avg": round(r.avg), "p95": round(r.p95)}
                                  for r in self.conc.itertuples()],
            "slope_ms_per_vu": round(self.slope, 1), "korelasi_vu_vs_avg": round(self.corr, 2),
            "endpoint": ep,
            "catatan_kualitas_data": self.dq,
            "temuan_otomatis": [{"severity": f["severity"], "judul": f["title"], "bukti": f["evidence"]} for f in self.findings],
        }
        persona = self.params.text("Peran dan Gaya Laporan")
        rules = "\n".join(f"- {r}" for r in self.params.items("Aturan Analisis"))
        ctx = "\n".join(f"- {k}: {v}" for k, v in self.context.items()) or "- (tidak diisi)"
        fmt = self.params.text("Format Output Rekomendasi")
        return f"""{persona}

KONTEKS SISTEM:
{ctx}

DATA HASIL PENGUJIAN (JSON):
{json.dumps(payload, indent=2, ensure_ascii=False, default=float)}

ATURAN ANALISIS:
{rules}

FORMAT OUTPUT:
{fmt}

Tulis dalam {self.llm['language']}. Berikan {self.llm['min_points']} sampai {self.llm['max_points']} poin utama,
diurutkan dari keparahan tertinggi, dan gunakan angka dari data di atas saja. Jangan mengarang angka."""

    def llm_recommendations(self, provider="groq"):
        if not self.llm.get("enabled"):
            return None
            
        print(f"[*] Menghubungi LLM via {provider.upper()} ({self.llm['model']}) ...")
        try:
            # --- OPSI 1: GROQ API ---
            if provider.lower() == "groq":
                api_key = self.llm.get("api_key", "")
                # Ambil key dari Streamlit Secrets jika ada
                try:
                    import streamlit as st
                    if "GROQ_API_KEY" in st.secrets:
                        api_key = st.secrets["GROQ_API_KEY"]
                except Exception:
                    pass

                headers = {
                    "Content-Type": "application/json",
                    "Authorization": f"Bearer {api_key}"
                }

                model_name = self.llm.get("model", "llama-3.1-8b-instant")
                
                payload = {
                    "model": self.llm.get("model", "llama-3.3-70b-versatile"),
                    "temperature": float(self.llm["temperature"]),
                    "messages": [{"role": "user", "content": self._llm_prompt()}]
                }
                
                url = self.llm.get("groq_api_url", "https://api.groq.com/openai/v1/chat/completions")
                res = requests.post(url, headers=headers, timeout=int(self.llm["timeout_seconds"]), json=payload)
                
                if res.status_code == 200:
                    res_data = res.json()
                    content = res_data["choices"][0]["message"]["content"].strip()
                    content = re.sub(r"^```(?:markdown|md)?\s*|\s*```$", "", content, flags=re.M).strip()
                    print("[ok] Rekomendasi Groq LLM diterima.")
                    return content
                else:
                    error_msg = f"HTTP {res.status_code}: {res.text}"
                    print(f"[!] Groq Error: {error_msg}")
                    return f"ERROR_GROQ: {error_msg}"

            # --- OPSI 2: OLLAMA LOKAL ---
            else:
                headers = {"Content-Type": "application/json"}
                url = self.llm.get("api_url", "http://localhost:11434/api/generate")
                
                if "/api/generate" in url:
                    payload = {
                        "model": "llama3.2",
                        "prompt": self._llm_prompt(),
                        "stream": False,
                        "options": {"temperature": float(self.llm["temperature"])}
                    }
                else:
                    payload = {
                        "model": "llama3.2",
                        "temperature": float(self.llm["temperature"]),
                        "messages": [{"role": "user", "content": self._llm_prompt()}]
                    }
                    
                res = requests.post(url, headers=headers, timeout=int(self.llm["timeout_seconds"]), json=payload)
                if res.status_code == 200:
                    res_data = res.json()
                    content = res_data.get("response", res_data.get("choices", [{}])[0].get("message", {}).get("content", "")).strip()
                    content = re.sub(r"^```(?:markdown|md)?\s*|\s*```$", "", content, flags=re.M).strip()
                    print("[ok] Rekomendasi Ollama diterima.")
                    return content

        except Exception as exc:  # noqa: BLE001
            print(f"[!] Gagal memanggil LLM ({provider}): {exc}; memakai playbook.")
            
        return None

    # ---- charts ----------------------------------------------------------- #
    def _charts(self, outdir):
        paths = {}
        labels = list(self.eps)
        short = [l if len(l) <= 30 else l[:29] + "…" for l in labels]
        p95 = [self.eps[l]["p95"] for l in labels]
        err = [self.eps[l]["error_rate"] for l in labels]

        fig, (a1, a2) = plt.subplots(1, 2, figsize=(11, max(3.6, 0.38 * len(labels) + 1.2)))
        b1 = a1.barh(short, p95, color="#1A365D")
        a1.axvline(self.sla["max_p95_ms"], color="red", ls="--", label=f"SLA {self.sla['max_p95_ms']} ms")
        a1.set_title("P95 Response Time (ms)", fontsize=10, fontweight="bold")
        a1.invert_yaxis(); a1.legend(loc="lower right", fontsize=8)
        for b, v in zip(b1, p95):
            if v > self.sla["max_p95_ms"]:
                b.set_color("#C53030")
        b2 = a2.barh(short, err, color="#2B6CB0")
        a2.axvline(self.sla["max_error_rate_pct"], color="red", ls="--", label=f"SLA {self.sla['max_error_rate_pct']}%")
        a2.set_title("Error Rate (%)", fontsize=10, fontweight="bold")
        a2.invert_yaxis(); a2.set_yticklabels([]); a2.legend(loc="lower right", fontsize=8)
        for b, v in zip(b2, err):
            if v > self.sla["max_error_rate_pct"]:
                b.set_color("#C53030")
        plt.tight_layout()
        paths["endpoint"] = os.path.join(outdir, "endpoint.png")
        plt.savefig(paths["endpoint"], dpi=200); plt.close()

        if not self.conc.empty:
            fig, ax = plt.subplots(figsize=(8, 3.6))
            ax.plot(self.conc["vu"], self.conc["avg"], marker="o", color="#1A365D", label="Avg")
            ax.plot(self.conc["vu"], self.conc["p95"], marker="s", color="#C53030", label="P95")
            ax.axhline(self.sla["max_p95_ms"], color="gray", ls="--", label=f"SLA P95 {self.sla['max_p95_ms']} ms")
            ax.set_xlabel("Virtual user aktif (allThreads)"); ax.set_ylabel("Response time (ms)")
            ax.set_title("Response Time vs Konkurensi", fontsize=10, fontweight="bold")
            ax.legend(fontsize=8); ax.grid(alpha=.3)
            plt.tight_layout()
            paths["concurrency"] = os.path.join(outdir, "concurrency.png")
            plt.savefig(paths["concurrency"], dpi=200); plt.close()
        return paths

    # ---- helper docx ------------------------------------------------------ #
    @staticmethod
    def _shade(cell, hex_color):
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>'))

    @staticmethod
    def _repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)

    @staticmethod
    def _rich(par, text, size=None, color=None):
        for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if not chunk:
                continue
            r = par.add_run(chunk)
            r.bold = bool(i % 2)
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color

    def _table(self, doc, headers, rows, widths, font=8.5, fills=None, align_right=()):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]; c.width = Inches(widths[i]); c.text = ""
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = WHITE
            self._shade(c, "1A365D")
        self._repeat_header(t.rows[0])
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].width = Inches(widths[i]); cells[i].text = ""
                p = cells[i].paragraphs[0]
                self._rich(p, str(val), size=font)
                if i in align_right:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if ri % 2:
                    self._shade(cells[i], "F7FAFC")
            if fills:
                for col, color in (fills(ri, row) or {}).items():
                    self._shade(cells[col], color)
        return t

    def _heading(self, doc, text, level=1):
        h = doc.add_heading(level=level)
        r = h.add_run(text); r.font.color.rgb = NAVY if level == 1 else DARK_GRAY; r.font.name = "Calibri"
        return h

    @staticmethod
    def _page_number(par):
        run = par.add_run()
        for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if typ:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), typ)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
            run._r.append(el)

    def _render_markdown(self, doc, text):
        """Render teks markdown sederhana dari LLM ke paragraf Word."""
        for raw in text.splitlines():
            if not raw.strip():
                continue
            indent = len(raw) - len(raw.lstrip(" "))
            s = raw.strip()
            m_head = re.match(r"^(?:#{1,4}\s+|\d+\.\s+)(.*)$", s)
            m_bul = re.match(r"^[*\-•]\s+(.*)$", s)
            if m_head and not m_bul and indent < 2:
                title = re.sub(r"^\*\*(.+?)\*\*$", r"\1", m_head.group(1).strip())
                self._heading(doc, title, level=3)
            elif m_bul:
                p = doc.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
                self._rich(p, m_bul.group(1), size=10)
            else:
                self._rich(doc.add_paragraph(), s, size=10)

    # ---- dokumen ---------------------------------------------------------- #
    def export_word_report(self, output_filename):
        if self.df is None:
            self.load_data()
        self.compute()
        self.detect_findings()

        tmp = tempfile.mkdtemp(prefix="jmeter_charts_")
        try:
            charts = self._charts(tmp)
            llm_text = self.llm_recommendations()
            blocks = self.playbook_blocks()
            self._build_doc(output_filename, charts, llm_text, blocks)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        print(f"[ok] Laporan tersimpan: {output_filename}")

    def _build_doc(self, out, charts, llm_text, blocks):
        meta, g, sla = self.params.meta, self.g, self.sla
        doc = Document()
        for s in doc.sections:
            s.top_margin = s.bottom_margin = Inches(0.9)
            s.left_margin = s.right_margin = Inches(0.9)
        st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)

        sec = doc.sections[0]
        hp = sec.header.paragraphs[0]
        hp.text = f"{meta.get('project_name', 'Proyek')}  |  {meta.get('classification', 'Internal')}"
        hp.runs[0].font.size = Pt(8); hp.runs[0].font.color.rgb = DARK_GRAY
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Halaman "); r.font.size = Pt(8)
        self._page_number(fp)

        p95_ok = g["p95"] <= sla["max_p95_ms"]
        err_ok = g["error_rate"] <= sla["max_error_rate_pct"]
        tps_ok = g["tps"] >= sla["min_tps"]
        passed = p95_ok and err_ok and tps_ok

        # Judul
        p = doc.add_paragraph(); r = p.add_run(str(meta.get("report_title", "Laporan Analisis Performa Sistem")))
        r.font.size = Pt(24); r.bold = True; r.font.color.rgb = NAVY
        d = datetime.now()
        p = doc.add_paragraph()
        r = p.add_run(f"{meta.get('project_name', '')} - {meta.get('test_type', 'Load Test')} ({meta.get('environment', '')})")
        r.font.size = Pt(12); r.font.color.rgb = DARK_GRAY
        meta_rows = [("Tanggal pengujian", str(meta.get("test_date", "-"))),
                     ("Tanggal laporan", f"{d.day} {MONTHS_ID[d.month - 1]} {d.year}"),
                     ("Disusun oleh", str(meta.get("prepared_by", "-"))), ("Versi", str(meta.get("version", "-"))),
                     ("File sumber", os.path.basename(self.jtl_filepath))]
        self._table(doc, ["Informasi", "Keterangan"], meta_rows, [1.8, 4.9], font=9)
        doc.add_paragraph()

        # Status
        t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0); self._shade(c, "E6F4EA" if passed else "FDECEA")
        pr = c.paragraphs[0]; pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pr.add_run("STATUS PENGUJIAN: " + ("LULUS (MEMENUHI SLA)" if passed else "TIDAK LULUS (ADA PELANGGARAN SLA)"))
        rr.bold = True; rr.font.size = Pt(13); rr.font.color.rgb = GREEN if passed else RED
        doc.add_paragraph()

        # 1 Ringkasan eksekutif
        self._heading(doc, "1. Ringkasan Eksekutif")
        wall_h = g["wall_sec"] / 3600
        p = doc.add_paragraph()
        self._rich(p, f"Pengujian mengumpulkan **{g['total_req']:,} sampel** pada **{len(self.eps)} transaksi** dengan konkurensi "
                      f"maksimum **{int(g['max_threads']) if g['max_threads'] == g['max_threads'] else '-'} virtual user**. "
                      f"Rentang waktu data {wall_h:.1f} jam, dengan trafik aktif pada {g['active_sec']:,} detik. "
                      f"Berdasarkan SLA (P95 ≤ {sla['max_p95_ms']} ms, error ≤ {sla['max_error_rate_pct']}%, TPS ≥ {sla['min_tps']}), "
                      f"hasil pengujian **{'memenuhi' if passed else 'belum memenuhi'}** SLA.")
        rows = [
            ("P95 Response Time", f"{g['p95']:.0f} ms", f"≤ {sla['max_p95_ms']} ms", "PASS" if p95_ok else "FAIL"),
            ("Error Rate", f"{g['error_rate']:.2f}%", f"≤ {sla['max_error_rate_pct']}%", "PASS" if err_ok else "FAIL"),
            ("Throughput aktif (TPS)", f"{g['tps']:.2f} req/s", f"≥ {sla['min_tps']} req/s", "PASS" if tps_ok else "FAIL"),
            ("Average Response Time", f"{g['avg']:.0f} ms", "-", "-"),
            ("P50 / P90 / P99", f"{g['p50']:.0f} / {g['p90']:.0f} / {g['p99']:.0f} ms", "-", "-"),
            ("Puncak throughput", f"{g['peak_tps']:.0f} req/s", "-", "-"),
            ("Porsi Latency terhadap Elapsed", f"{g['latency_share']:.0f}%" if g["latency_share"] == g["latency_share"] else "-", "-", "-"),
            ("Total data diunduh", f"{g['total_bytes'] / 1e6:.0f} MB" if g["total_bytes"] == g["total_bytes"] else "-", "-", "-"),
        ]
        self._table(doc, ["Metrik", "Hasil", "Target SLA", "Status"], rows, [2.5, 1.9, 1.4, 0.9], font=9,
                    fills=lambda i, r: {3: "C6F6D5" if r[3] == "PASS" else "FED7D7"} if r[3] in ("PASS", "FAIL") else None)
        doc.add_paragraph()

        # 2 Temuan utama
        self._heading(doc, "2. Temuan Utama")
        if self.findings:
            rows = [(i + 1, f["severity"], f["title"], f["evidence"]) for i, f in enumerate(self.findings)]
            self._table(doc, ["#", "Severity", "Temuan", "Bukti dari data"], rows, [0.3, 0.8, 1.9, 3.7], font=8.5,
                        fills=lambda i, r: {1: SEVERITY_FILL[r[1]]})
        else:
            doc.add_paragraph("Tidak ada temuan otomatis. Seluruh metrik berada dalam ambang yang ditetapkan.")
        doc.add_paragraph()

        # 3 Dashboard
        self._heading(doc, "3. Dashboard Visual")
        doc.add_picture(charts["endpoint"], width=Inches(6.6)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if "concurrency" in charts:
            doc.add_picture(charts["concurrency"], width=Inches(5.6)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            self._rich(cap, f"Slope kenaikan response time **{self.slope:.0f} ms per virtual user** (korelasi {self.corr:.2f}).", size=9)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 4 Ringkasan semua API
        self._heading(doc, "4. Ringkasan Seluruh Endpoint")
        rows = []
        for l, s in sorted(self.eps.items(), key=lambda x: -x[1]["p95"]):
            ok = s["p95"] <= sla["max_p95_ms"] and s["error_rate"] <= sla["max_error_rate_pct"]
            kb = f"{s['avg_bytes'] / 1024:.1f}" if s["avg_bytes"] == s["avg_bytes"] else "-"
            rows.append((l, f"{s['total_req']:,}", f"{s['avg']:.0f}", f"{s['p95']:.0f}", f"{s['p99']:.0f}",
                         f"{s['error_rate']:.2f}%", kb, "PASS" if ok else "FAIL"))
        self._table(doc, ["Endpoint (urut P95)", "Req", "Avg", "P95", "P99", "Error", "KB", "SLA"], rows,
                    [2.4, 0.55, 0.5, 0.5, 0.5, 0.6, 0.65, 0.5], font=8, align_right=(1, 2, 3, 4, 5, 6),
                    fills=lambda i, r: {7: "C6F6D5" if r[7] == "PASS" else "FED7D7"})
        doc.add_paragraph()

        # 5 Fase dan konkurensi
        self._heading(doc, "5. Analisis Fase Pengujian dan Konkurensi")
        if self.phases:
            rows = [(k, f"{v['total_req']:,}", f"{v['avg']:.0f}", f"{v['p95']:.0f}", f"{v['p99']:.0f}", f"{v['tps']:.2f}",
                     f"{v['error_rate']:.2f}%") for k, v in self.phases.items()]
            self._table(doc, ["Fase", "Req", "Avg (ms)", "P95 (ms)", "P99 (ms)", "TPS aktif", "Error"], rows,
                        [1.8, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8], font=9, align_right=(1, 2, 3, 4, 5, 6))
            doc.add_paragraph()
        if not self.conc.empty:
            rows = [(int(r.vu), f"{int(r.samples):,}", f"{r.avg:.0f}", f"{r.p95:.0f}") for r in self.conc.itertuples()]
            self._table(doc, ["Virtual user aktif", "Sampel", "Avg (ms)", "P95 (ms)"], rows, [1.8, 1.2, 1.2, 1.2], font=8.5,
                        align_right=(1, 2, 3))
            doc.add_paragraph()

        # 6 Detail per endpoint
        self._heading(doc, "6. Observasi per Endpoint")
        for l, s in self.eps.items():
            self._heading(doc, l, level=2)
            self._table(doc, ["Req", "Error", "Avg", "P50", "P95", "P99", "Latency", "Connect"],
                        [(f"{s['total_req']:,}", f"{s['error_rate']:.2f}%", f"{s['avg']:.0f}", f"{s['p50']:.0f}", f"{s['p95']:.0f}",
                          f"{s['p99']:.0f}", f"{s['avg_latency']:.0f}", f"{s['avg_connect']:.1f}")],
                        [0.75] * 8, font=8.5, align_right=tuple(range(8)))
            for ins in self.endpoint_insights(s):
                self._rich(doc.add_paragraph(style="List Bullet"), ins, size=9.5)

        # 7 Rekomendasi
        self._heading(doc, "7. Rekomendasi Teknis dan Arsitektur")
        if llm_text:
            n = self.llm["model"]
            self._rich(doc.add_paragraph(), f"Analisis disusun oleh model **{n}** berdasarkan data pengujian dan konteks sistem pada file parameter.", size=9)
            self._render_markdown(doc, llm_text)
        elif blocks:
            for i, b in enumerate(blocks, 1):
                self._heading(doc, f"{i}. {b['judul']}  [{b['severity']}]", level=3)
                for lab, key in (("Temuan Metrik", "evidence"), ("Root Cause Analysis", "rca"), ("Rekomendasi Database", "database"),
                                 ("Rekomendasi Application / Connection Pool", "application"),
                                 ("Rekomendasi Infrastructure", "infrastructure"), ("Verifikasi", "verifikasi")):
                    if b.get(key):
                        self._rich(doc.add_paragraph(style="List Bullet"), f"**{lab}:** {b[key]}", size=10)
        else:
            doc.add_paragraph("Tidak ada rekomendasi karena tidak ada temuan yang terdeteksi.")

        # 8 Matriks tindakan
        if blocks:
            rows = [(i, b["severity"], b["judul"], b["database"], b["application"], b["infrastructure"])
                    for i, b in enumerate(blocks, 1)]
            sec2 = doc.add_section()
            from docx.enum.section import WD_ORIENT
            sec2.orientation = WD_ORIENT.LANDSCAPE
            sec2.page_width, sec2.page_height = sec2.page_height, sec2.page_width
            sec2.left_margin = sec2.right_margin = Inches(0.6)
            self._heading(doc, "8. Matriks Tindakan Prioritas")
            self._table(doc, ["#", "Severity", "Temuan", "Database", "Application / Pool", "Infrastructure"], rows,
                        [0.3, 0.7, 1.6, 2.4, 2.4, 2.4], font=7.5, fills=lambda i, r: {1: SEVERITY_FILL[r[1]]})

        # Lampiran
        self._heading(doc, "Lampiran A. Kualitas Data dan Parameter")
        for note in self.dq:
            self._rich(doc.add_paragraph(style="List Bullet"), note, size=9.5)
        if self.context:
            self._rich(doc.add_paragraph(), "**Konteks sistem yang dipakai:**", size=9.5)
            self._table(doc, ["Komponen", "Keterangan"], [(k, v) for k, v in self.context.items()], [1.8, 4.9], font=8.5)
        for note in self.params.items("Catatan Tambahan"):
            self._rich(doc.add_paragraph(style="List Bullet"), note, size=9)

        doc.save(out)


# --------------------------------------------------------------------------- #
# 3. CLI
# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description="Generator laporan Word dari hasil JMeter dengan parameter Markdown.")
    ap.add_argument("--jtl", default="result.jtl", help="Path file hasil JMeter (.jtl / CSV)")
    ap.add_argument("--params", default="report_parameters.md", help="Path file parameter Markdown")
    ap.add_argument("--output", default="Laporan_Lengkap_JMeter_Performance.docx", help="Nama file Word output")
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--llm", dest="llm", action="store_true", default=None, help="Paksa LLM aktif")
    g.add_argument("--no-llm", dest="llm", action="store_false", help="Nonaktifkan LLM (pakai playbook)")
    a, unknown = ap.parse_known_args()

    params = ReportParameters(a.params)
    JMeterReportGenerator(a.jtl, params, llm_override=a.llm).export_word_report(a.output)


if __name__ == "__main__":
    try:
        main()
    except Exception as err:  # noqa: BLE001
        print(f"[x] Error: {err}")
        raise SystemExit(1)